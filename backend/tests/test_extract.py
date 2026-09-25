import io

import pytest
from docx import Document
from pypdf import PdfWriter

from app.ingest.extract import extract_text


def test_markdown_and_txt_are_decoded():
    assert extract_text("a.md", b"# Hi\nthere") == "# Hi\nthere"
    assert extract_text("A.TXT", b"plain") == "plain"


def test_non_utf8_text_is_rejected():
    with pytest.raises(ValueError, match="UTF-8"):
        extract_text("a.txt", b"\xff\xfe\x00bad")


def test_unsupported_extension_is_rejected():
    with pytest.raises(ValueError, match="unsupported"):
        extract_text("a.exe", b"x")


def test_docx_headings_become_markdown():
    doc = Document()
    doc.add_heading("Refunds", level=1)
    doc.add_paragraph("Within 30 days.")
    buf = io.BytesIO()
    doc.save(buf)
    text = extract_text("policy.docx", buf.getvalue())
    assert text == "# Refunds\n\nWithin 30 days."


def test_pdf_without_text_is_rejected():
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buf = io.BytesIO()
    writer.write(buf)
    with pytest.raises(ValueError, match="no extractable text"):
        extract_text("scan.pdf", buf.getvalue())


def test_corrupt_pdf_is_rejected():
    with pytest.raises(ValueError, match="could not read"):
        extract_text("bad.pdf", b"not a pdf")
